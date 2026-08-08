import os
import json
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, FileResponse, HttpResponse
from django.views.decorators.http import require_POST
from django.utils import timezone
from .models import Papka, Fayl


def hajm_formatlash(hajm):
    for birlik in ["B", "KB", "MB", "GB"]:
        if hajm < 1024:
            return f"{hajm:.1f} {birlik}"
        hajm /= 1024
    return f"{hajm:.1f} TB"


@login_required
def dashboard(request, papka_id=None):
    joriy_papka = None
    if papka_id:
        joriy_papka = get_object_or_404(Papka, id=papka_id, foydalanuvchi=request.user)

    papkalar = Papka.objects.filter(foydalanuvchi=request.user, ota_papka=joriy_papka).order_by("nom")
    fayllar = Fayl.objects.filter(foydalanuvchi=request.user, papka=joriy_papka).order_by("-yangilangan")

    non = joriy_papka.non_olish() if joriy_papka else []

    # Umumiy hajm
    jami_hajm = sum(f.hajm for f in Fayl.objects.filter(foydalanuvchi=request.user))
    fayllar_soni = Fayl.objects.filter(foydalanuvchi=request.user).count()

    context = {
        "joriy_papka": joriy_papka,
        "papkalar": papkalar,
        "fayllar": fayllar,
        "non": non,
        "jami_hajm": hajm_formatlash(jami_hajm),
        "fayllar_soni": fayllar_soni,
    }
    return render(request, "storage/dashboard.html", context)


@login_required
def fayl_yuklash(request):
    if request.method == "POST":
        papka_id = request.POST.get("papka_id") or None
        papka = None
        if papka_id:
            papka = get_object_or_404(Papka, id=papka_id, foydalanuvchi=request.user)

        yuklangan_fayllar = request.FILES.getlist("fayllar")
        natijalar = []

        for yuklangan in yuklangan_fayllar:
            fayl = Fayl.objects.create(
                foydalanuvchi=request.user,
                papka=papka,
                nom=yuklangan.name,
                fayl=yuklangan,
                hajm=yuklangan.size,
            )
            natijalar.append({
                "id": fayl.id,
                "nom": fayl.nom,
                "ikonka": fayl.ikonka(),
                "hajm": fayl.hajm_formatlangan(),
                "tur": fayl.tur(),
                "rang": fayl.rang_klass(),
            })

        return JsonResponse({"muvaffaqiyat": True, "fayllar": natijalar})
    return JsonResponse({"xato": "Noto'g'ri so'rov"}, status=400)


@login_required
def papka_yaratish(request):
    if request.method == "POST":
        try:
            data = json.loads(request.body)
        except json.JSONDecodeError:
            data = request.POST

        nom = (data.get("nom") or "").strip()
        papka_id = data.get("papka_id") or None

        if not nom:
            return JsonResponse({"xato": "Papka nomi kiritilishi shart"}, status=400)

        ota = None
        if papka_id:
            ota = get_object_or_404(Papka, id=papka_id, foydalanuvchi=request.user)

        papka = Papka.objects.create(foydalanuvchi=request.user, nom=nom, ota_papka=ota)
        return JsonResponse({"muvaffaqiyat": True, "id": papka.id, "nom": papka.nom})

    return JsonResponse({"xato": "Noto'g'ri so'rov"}, status=400)


@login_required
def fayl_korish(request, fayl_id):
    fayl = get_object_or_404(Fayl, id=fayl_id, foydalanuvchi=request.user)

    context = {"fayl": fayl, "korsatish_turi": "yuklab_olish", "kontent": None}

    if fayl.matn_mi():
        try:
            with open(fayl.fayl.path, "r", encoding="utf-8", errors="replace") as f:
                kontent = f.read()
            context["kontent"] = kontent
            context["korsatish_turi"] = "kod"
        except Exception as e:
            context["xato"] = str(e)

    elif fayl.word_mi():
        try:
            from docx import Document
            doc = Document(fayl.fayl.path)
            html_qismlar = []
            for para in doc.paragraphs:
                text = para.text
                if not text.strip():
                    html_qismlar.append("<p><br></p>")
                    continue
                stil = para.style.name
                if "Heading 1" in stil:
                    html_qismlar.append(f"<h1>{text}</h1>")
                elif "Heading 2" in stil:
                    html_qismlar.append(f"<h2>{text}</h2>")
                elif "Heading 3" in stil:
                    html_qismlar.append(f"<h3>{text}</h3>")
                else:
                    html_qismlar.append(f"<p>{text}</p>")
            context["kontent"] = "\n".join(html_qismlar)
            context["korsatish_turi"] = "word"
        except Exception as e:
            context["xato"] = f"Faylni o'qishda xato: {e}"
            context["korsatish_turi"] = "yuklab_olish"

    elif fayl.rasm_mi():
        context["korsatish_turi"] = "rasm"

    elif fayl.pdf_mi():
        context["korsatish_turi"] = "pdf"

    elif fayl.video_mi():
        context["korsatish_turi"] = "video"

    elif fayl.audio_mi():
        context["korsatish_turi"] = "audio"

    return render(request, "storage/viewer.html", context)


@login_required
def fayl_saqlash(request, fayl_id):
    if request.method != "POST":
        return JsonResponse({"xato": "Noto'g'ri so'rov"}, status=405)

    fayl = get_object_or_404(Fayl, id=fayl_id, foydalanuvchi=request.user)

    try:
        data = json.loads(request.body)
        kontent = data.get("kontent", "")
        pozitsiya = data.get("pozitsiya", {})

        if fayl.matn_mi():
            with open(fayl.fayl.path, "w", encoding="utf-8") as f:
                f.write(kontent)

        elif fayl.word_mi():
            import re
            from docx import Document
            doc = Document()

            # HTML ni oddiy matnlarga bo'lish
            kontent = re.sub(r"<h1[^>]*>", "\n__H1__", kontent)
            kontent = re.sub(r"</h1>", "", kontent)
            kontent = re.sub(r"<h2[^>]*>", "\n__H2__", kontent)
            kontent = re.sub(r"</h2>", "", kontent)
            kontent = re.sub(r"<h3[^>]*>", "\n__H3__", kontent)
            kontent = re.sub(r"</h3>", "", kontent)
            kontent = re.sub(r"<br\s*/?>", "\n", kontent)
            kontent = re.sub(r"<[^>]+>", "", kontent)
            kontent = kontent.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")

            for qator in kontent.split("\n"):
                qator = qator.strip()
                if not qator:
                    continue
                if qator.startswith("__H1__"):
                    doc.add_heading(qator[6:], level=1)
                elif qator.startswith("__H2__"):
                    doc.add_heading(qator[6:], level=2)
                elif qator.startswith("__H3__"):
                    doc.add_heading(qator[6:], level=3)
                else:
                    doc.add_paragraph(qator)

            doc.save(fayl.fayl.path)

        fayl.hajm = os.path.getsize(fayl.fayl.path)
        fayl.oxirgi_pozitsiya = pozitsiya
        fayl.save()

        return JsonResponse({"muvaffaqiyat": True})

    except Exception as e:
        return JsonResponse({"xato": str(e)}, status=500)


@login_required
@require_POST
def fayl_ochirish(request, fayl_id):
    fayl = get_object_or_404(Fayl, id=fayl_id, foydalanuvchi=request.user)
    try:
        os.remove(fayl.fayl.path)
    except Exception:
        pass
    fayl.delete()
    return JsonResponse({"muvaffaqiyat": True})


@login_required
@require_POST
def papka_ochirish(request, papka_id):
    papka = get_object_or_404(Papka, id=papka_id, foydalanuvchi=request.user)

    def rekursiv_ochirish(p):
        for f in p.fayllar.all():
            try:
                os.remove(f.fayl.path)
            except Exception:
                pass
            f.delete()
        for ichki in p.ichki_papkalar.all():
            rekursiv_ochirish(ichki)
        p.delete()

    rekursiv_ochirish(papka)
    return JsonResponse({"muvaffaqiyat": True})


@login_required
def fayl_yuklab_olish(request, fayl_id):
    fayl = get_object_or_404(Fayl, id=fayl_id, foydalanuvchi=request.user)
    response = FileResponse(open(fayl.fayl.path, "rb"))
    response["Content-Disposition"] = f'attachment; filename="{fayl.nom}"'
    return response


@login_required
def fayl_nomini_ozgartirish(request, fayl_id):
    if request.method == "POST":
        fayl = get_object_or_404(Fayl, id=fayl_id, foydalanuvchi=request.user)
        try:
            data = json.loads(request.body)
            yangi_nom = (data.get("nom") or "").strip()
            if yangi_nom:
                # Kengaytmani saqlash
                eski_kengaytma = os.path.splitext(fayl.nom)[1]
                yangi_kengaytma = os.path.splitext(yangi_nom)[1]
                if not yangi_kengaytma:
                    yangi_nom += eski_kengaytma
                fayl.nom = yangi_nom
                fayl.save()
                return JsonResponse({"muvaffaqiyat": True, "nom": yangi_nom})
        except Exception as e:
            return JsonResponse({"xato": str(e)}, status=400)
    return JsonResponse({"xato": "Noto'g'ri so'rov"}, status=405)
